from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from build_rag_meeting_ppt import (
    ROOT,
    ASSETS,
    WIDE_W,
    WIDE_H,
    BG,
    INK,
    MUTED,
    BLUE,
    TEAL,
    AMBER,
    RED,
    GREEN,
    LINE,
    WHITE,
    DARK,
    add_bg,
    add_title,
    add_footer,
    add_badge,
    add_card,
    add_bullets,
    add_flow,
    add_image_fit,
    add_section_label,
    set_fill,
)


OUT = ROOT / "docs" / "RA-RAG-组会汇报-详细版.pptx"


def add_table(slide, rows, cols, x, y, w, h, data, header_fill=RGBColor(226, 232, 240)):
    table_shape = slide.shapes.add_table(rows, cols, x, y, w, h)
    table = table_shape.table
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = data[r][c]
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            fill = cell.fill
            fill.solid()
            fill.fore_color.rgb = header_fill if r == 0 else WHITE
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT
                for run in p.runs:
                    run.font.name = "Microsoft YaHei"
                    run.font.size = Pt(9.3 if rows >= 6 else 10.5)
                    run.font.bold = r == 0
                    run.font.color.rgb = INK if r == 0 else MUTED
    return table_shape


def add_code_box(slide, text, x, y, w, h, title="代码锚点"):
    box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(15, 23, 42)
    box.line.color.rgb = RGBColor(51, 65, 85)
    title_box = slide.shapes.add_textbox(x + Inches(0.18), y + Inches(0.12), w - Inches(0.35), Inches(0.25))
    title_box.text_frame.text = title
    p = title_box.text_frame.paragraphs[0]
    p.runs[0].font.name = "Microsoft YaHei"
    p.runs[0].font.size = Pt(9.5)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor(147, 197, 253)
    body = slide.shapes.add_textbox(x + Inches(0.18), y + Inches(0.42), w - Inches(0.35), h - Inches(0.48))
    body.text_frame.text = text
    body.text_frame.word_wrap = True
    for p in body.text_frame.paragraphs:
        for run in p.runs:
            run.font.name = "Consolas"
            run.font.size = Pt(9.8)
            run.font.color.rgb = RGBColor(226, 232, 240)
    return box


def add_big_number(slide, number, label, body, x, y, accent):
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, Inches(3.75), Inches(1.45))
    set_fill(card, WHITE)
    card.line.color.rgb = LINE
    num = slide.shapes.add_textbox(x + Inches(0.18), y + Inches(0.16), Inches(0.75), Inches(0.48))
    num.text_frame.text = str(number)
    p = num.text_frame.paragraphs[0]
    p.runs[0].font.name = "Microsoft YaHei"
    p.runs[0].font.size = Pt(24)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = accent
    lb = slide.shapes.add_textbox(x + Inches(0.95), y + Inches(0.16), Inches(2.6), Inches(0.32))
    lb.text_frame.text = label
    p = lb.text_frame.paragraphs[0]
    p.runs[0].font.name = "Microsoft YaHei"
    p.runs[0].font.size = Pt(14)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = INK
    bd = slide.shapes.add_textbox(x + Inches(0.95), y + Inches(0.55), Inches(2.55), Inches(0.68))
    bd.text_frame.text = body
    for p in bd.text_frame.paragraphs:
        for run in p.runs:
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(10.8)
            run.font.color.rgb = MUTED


def make_detailed_deck():
    prs = Presentation()
    prs.slide_width = WIDE_W
    prs.slide_height = WIDE_H
    blank = prs.slide_layouts[6]

    # 1
    slide = prs.slides.add_slide(blank)
    add_bg(slide, DARK)
    slide.shapes.add_picture(str(ASSETS / "ragent-ai-banner.png"), Inches(8.0), Inches(0.4), width=Inches(4.7))
    add_badge(slide, "组会分享", Inches(0.75), Inches(0.82), TEAL)
    title = slide.shapes.add_textbox(Inches(0.75), Inches(1.46), Inches(8.85), Inches(1.55))
    title.text_frame.text = "从 RA 项目看 RAG 系统的工程化实践"
    p = title.text_frame.paragraphs[0]
    p.runs[0].font.name = "Microsoft YaHei"
    p.runs[0].font.size = Pt(35)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = WHITE
    sub = slide.shapes.add_textbox(Inches(0.78), Inches(3.05), Inches(9.2), Inches(0.75))
    sub.text_frame.text = "更详细版：从文档入库、意图路由、多路召回、Rerank 到 Trace / 评估闭环"
    p = sub.text_frame.paragraphs[0]
    p.runs[0].font.name = "Microsoft YaHei"
    p.runs[0].font.size = Pt(17)
    p.runs[0].font.color.rgb = RGBColor(203, 213, 225)
    add_flow(slide, ["离线入库", "在线链路", "检索优化", "可观测评估"], Inches(0.82), Inches(5.2), Inches(2.15), Inches(0.35), color=TEAL)
    add_footer(slide, 1)

    # 2
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "这次汇报怎么组织", "把 RAG 主题重组为一条项目链路：先有知识，再有检索，最后有诊断", 2)
    add_big_number(slide, 1, "从知识入库开始", "解析、分块、向量化和索引，决定召回上限。", Inches(0.78), Inches(1.55), BLUE)
    add_big_number(slide, 2, "再讲在线 Pipeline", "记忆、改写、意图、短路、检索、生成。", Inches(4.82), Inches(1.55), TEAL)
    add_big_number(slide, 3, "重点讲检索优化", "意图定向 + 全局兜底 + 去重 + 重排。", Inches(8.86), Inches(1.55), AMBER)
    add_card(slide, "本版增加的细节", "关键配置：topK、阈值、multiplier、记忆窗口\n关键代码：Pipeline、IntentResolver、RetrievalEngine、PostProcessor\n关键诊断：Trace 节点、指标联读、问题归因路径", Inches(0.95), Inches(3.65), Inches(5.6), Inches(1.9), GREEN, 16, 13)
    add_card(slide, "仍然不展开的内容", "不讲模型训练细节、不讲向量数据库内核、不讲 MCP 协议专题、不讲前端实现。\n\n这些会占时间，但对本次组会主线帮助不大。", Inches(6.9), Inches(3.65), Inches(5.6), Inches(1.9), RED, 16, 13)

    # 3
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "RA 的 RAG 系统全景", "不是“问题 -> 向量库 -> LLM”，而是一条可短路、可扩展、可观测的链路", 3)
    add_image_fit(slide, ASSETS / "ragent-chain-v3.png", Inches(0.58), Inches(1.42), Inches(7.15), Inches(5.35))
    add_card(slide, "三个阶段", "1. 入库：文档解析、分块、Embedding、索引\n2. 检索：改写、意图路由、多通道召回、Rerank\n3. 生成：证据组装、模板选择、流式输出", Inches(8.0), Inches(1.55), Inches(4.75), Inches(1.9), BLUE, 15, 12.8)
    add_card(slide, "三个短路", "1. 歧义：先引导澄清\n2. SYSTEM：不检索直接回答\n3. 空召回：明确告知无相关文档", Inches(8.0), Inches(3.75), Inches(4.75), Inches(1.5), TEAL, 15, 12.8)
    add_code_box(slide, "StreamChatPipeline.execute(ctx)", Inches(8.0), Inches(5.55), Inches(4.75), Inches(0.75))

    # 4
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "为什么 RA 项目不能只靠大模型", "项目里的知识是私域、动态、可审计的；RAG 是把这些知识接入模型的工程通道", 4)
    data = [
        ["问题", "只问大模型", "RAG 方案"],
        ["私域知识", "模型不知道企业内部制度、业务系统规则", "从知识库检索真实文档作为证据"],
        ["知识更新", "模型参数不会随文档变化自动更新", "更新文档 / 重建索引即可生效"],
        ["可追溯", "答案来源不清，难复盘", "答案可回看召回 chunk、trace 节点和指标"],
        ["幻觉风险", "可能靠预训练知识补全细节", "用 Prompt + evidence 约束回答边界"],
    ]
    add_table(slide, 5, 3, Inches(0.75), Inches(1.55), Inches(11.85), Inches(3.1), data)
    add_card(slide, "讲述落点", "RAG 不是让模型“更聪明”，而是让模型在回答时有可控知识来源。\n\n这也是后面多路召回、Rerank、Trace、评估存在的原因。", Inches(1.2), Inches(5.25), Inches(10.9), Inches(1.05), GREEN, 15, 13)

    # 5
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "离线入库：从原始文档到向量索引", "入库管道是 RAG 的地基；在线检索只是使用已经加工好的知识", 5)
    add_image_fit(slide, ASSETS / "ingestion-pipeline.png", Inches(0.58), Inches(1.35), Inches(6.35), Inches(5.3))
    data = [
        ["节点", "职责", "产物"],
        ["Fetcher", "拉取本地、远程、S3、Feishu 等来源", "DocumentSource"],
        ["Parser", "按文件类型解析文本、表格、图片、PDF 结构", "ParsedDocument / Block"],
        ["Chunker", "结构化分块，生成可检索片段", "VectorChunk"],
        ["Enhancer / Enricher", "可选增强、补充摘要或说明", "增强后的 chunk 文本"],
        ["Indexer", "Embedding 后写入向量存储", "KnowledgeVector"],
    ]
    add_table(slide, 6, 3, Inches(7.15), Inches(1.45), Inches(5.45), Inches(3.3), data)
    add_card(slide, "细节提醒", "入库质量差时，在线阶段通常只能“更努力地找错东西”。所以文档解析、分块粒度和 metadata 比很多人想象得更关键。", Inches(7.15), Inches(5.15), Inches(5.45), Inches(0.95), AMBER, 14, 12)

    # 6
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "解析层：不同文档不是同一种处理方式", "Parser 的目标不是简单抽文本，而是尽量保留文档结构和可回答信息", 6)
    data = [
        ["类型", "项目解析器", "保留重点"],
        ["Markdown / Text", "MarkdownDocumentParser / Tika", "标题层级、段落、列表"],
        ["Excel / CSV", "ExcelDocumentParser / CsvDocumentParser", "表格结构、单元格值、超链接"],
        ["PDF / Word / PPT", "MinerU / Tika", "版面结构、表格、公式、图片资源"],
        ["Image", "ImageDocumentParser + VLM", "OCR 文本、图表关系、图片说明"],
        ["混合文档", "DocumentParserSelector", "按类型选择最合适的解析器"],
    ]
    add_table(slide, 6, 3, Inches(0.78), Inches(1.55), Inches(7.0), Inches(3.65), data)
    add_card(slide, "为什么要保留结构", "RAG 的召回单位不是“文件”，而是 Chunk。\n\n如果表格、标题、列表、图片说明在解析阶段丢了，后面 Embedding 和检索都无法恢复这些信息。", Inches(8.1), Inches(1.65), Inches(4.45), Inches(2.0), BLUE, 15, 13)
    add_card(slide, "项目素材可以展示", "可以用后台知识库 / 数据集截图说明：RA 不只是文本 QA，也在处理多类型知识资产。", Inches(8.1), Inches(4.15), Inches(4.45), Inches(1.25), TEAL, 15, 13)

    # 7
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Chunking 细节：语义边界、体量预算、特殊块", "结构化 Chunking 是提高 context recall 的关键，不只是调一个 chunkSize", 7)
    add_card(slide, "优先级 1：整篇不分块", "当 chunkSize / targetChars = -1 时，整篇合成一个 DOCUMENT chunk。适合短文档或必须整体保真的材料。", Inches(0.75), Inches(1.45), Inches(3.9), Inches(1.45), BLUE, 14, 11.7)
    add_card(slide, "优先级 2：Block-aware", "如果 Parser 产出结构化 blocks，就按 heading、paragraph、table、list、image、code 分派处理。", Inches(4.85), Inches(1.45), Inches(3.9), Inches(1.45), TEAL, 14, 11.7)
    add_card(slide, "优先级 3：Legacy 文本策略", "blocks 为空时回退 fixed-size / structure-aware text chunker，避免无结构文档无法入库。", Inches(8.95), Inches(1.45), Inches(3.9), Inches(1.45), AMBER, 14, 11.7)
    data = [
        ["参数", "默认值", "含义"],
        ["DEFAULT_MAX_CHARS", "512", "每个 Chunk 的体量预算"],
        ["DEFAULT_OVERLAP", "64", "段落重叠，缓解语义断裂"],
        ["DEFAULT_ROWS_PER_CHUNK", "50", "表格每块最大行数"],
        ["DEFAULT_LIST_ITEMS_PER_CHUNK", "10", "长列表分块 item 数"],
    ]
    add_table(slide, 5, 3, Inches(1.2), Inches(3.55), Inches(10.8), Inches(2.15), data)
    add_section_label(slide, "这页建议讲一个表格/列表切分例子", Inches(5.3), Inches(6.18), GREEN)

    # 8
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Embedding 与向量存储：讲工程参数，不讲算法推导", "这页只讲项目怎么接模型、怎么存向量、怎么保持可替换", 8)
    add_card(slide, "向量维度", "当前默认 dimension = 1536。\n\n所有 embedding 模型候选需要与向量库字段维度保持一致。", Inches(0.75), Inches(1.55), Inches(3.75), Inches(1.45), BLUE, 15, 12.5)
    add_card(slide, "相似度", "metric-type = COSINE。\n\nPgVector 创建 HNSW + vector_cosine_ops 索引。", Inches(4.78), Inches(1.55), Inches(3.75), Inches(1.45), TEAL, 15, 12.5)
    add_card(slide, "存储可切换", "rag.vector.type = pg。\n\n项目同时提供 PgVectorStoreService 和 MilvusVectorStoreService。", Inches(8.82), Inches(1.55), Inches(3.75), Inches(1.45), AMBER, 15, 12.5)
    data = [
        ["模型候选", "Provider", "角色"],
        ["qwen-emb-8b", "siliconflow", "默认 embedding"],
        ["qwen-emb-local", "ollama", "本地兜底"],
        ["text-embedding-3-large", "aihubmix", "外部候选"],
    ]
    add_table(slide, 4, 3, Inches(1.55), Inches(4.0), Inches(10.2), Inches(1.55), data)
    add_card(slide, "讲述重点", "这里不展开 embedding 算法，只说明：模型路由 + 向量存储抽象让 RA 可以替换模型和存储后端。", Inches(2.3), Inches(6.0), Inches(8.7), Inches(0.65), GREEN, 13, 11.5)

    # 9
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "在线 Pipeline：一次问题的真实执行顺序", "后面所有检索优化都挂在这条主链路上", 9)
    labels = ["loadMemory", "rewriteQuery", "resolveIntents", "handleGuidance", "retrieve", "streamRagResponse"]
    add_flow(slide, labels, Inches(0.55), Inches(1.45), Inches(1.85), Inches(0.18), color=BLUE)
    data = [
        ["阶段", "输入", "输出/作用"],
        ["loadMemory", "conversationId, userId, question", "history + 摘要"],
        ["rewriteQuery", "question + history", "RewriteResult + subQuestions"],
        ["resolveIntents", "subQuestions", "SubQuestionIntent 列表"],
        ["handleGuidance", "意图分数", "歧义时短路返回引导"],
        ["retrieve", "意图 + topK", "RetrievalContext"],
        ["streamRagResponse", "context + prompt plan", "LLM 流式回答"],
    ]
    add_table(slide, 7, 3, Inches(0.85), Inches(2.65), Inches(8.1), Inches(3.2), data)
    add_code_box(slide, "public void execute(StreamChatContext ctx) {\n  loadMemory(ctx);\n  rewriteQuery(ctx);\n  resolveIntents(ctx);\n  if (handleGuidance(ctx)) return;\n  if (handleSystemOnly(ctx)) return;\n  RetrievalContext rc = retrieve(ctx);\n  if (handleEmptyRetrieval(ctx, rc)) return;\n  streamRagResponse(ctx, rc);\n}", Inches(9.2), Inches(2.4), Inches(3.5), Inches(3.75), "StreamChatPipeline")

    # 10
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Memory + Query Rewrite：先补上下文，再改写问题", "长对话场景下，用户问题经常依赖历史，不能直接拿原句去检索", 10)
    add_card(slide, "记忆窗口", "history-keep-turns = 4\nsummary-start-turns = 5\nsummary-enabled = true\nsummary-max-chars = 200", Inches(0.75), Inches(1.55), Inches(3.7), Inches(1.65), BLUE, 15, 12.8)
    add_card(slide, "改写能力", "rewriteWithSplit(question, history)\n\n支持结合历史改写，并拆分多问句为多个子问题。", Inches(4.82), Inches(1.55), Inches(3.7), Inches(1.65), TEAL, 15, 12.8)
    add_card(slide, "失败兜底", "LLM 改写失败时回退原问题；拆分失败时仍保证至少一个查询可进入检索。", Inches(8.9), Inches(1.55), Inches(3.7), Inches(1.65), AMBER, 15, 12.8)
    add_card(slide, "例子", "用户连续问：\n第一轮：公司的报销流程是什么？\n第二轮：那发票抬头怎么开？\n\n第二轮需要结合历史，把“那”改写为“报销相关发票抬头/开票信息”。", Inches(1.35), Inches(4.05), Inches(10.7), Inches(1.55), GREEN, 16, 14)

    # 11
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "意图树：把问题路由到知识库、系统回复或工具", "意图不是为了分类好看，而是决定后续检索目标和 Prompt 场景", 11)
    data = [
        ["字段", "含义", "影响"],
        ["kind", "KB / MCP / SYSTEM", "决定检索、工具或直接回复"],
        ["level", "DOMAIN / CATEGORY / TOPIC", "组织意图层级"],
        ["kbId / collectionName", "绑定知识库和 collection", "定向检索目标"],
        ["mcpToolId", "绑定工具", "工具调用路由"],
        ["topK", "意图级召回数量", "覆盖全局 topK"],
        ["promptTemplate / snippet", "节点级 Prompt", "定制回答风格和规则"],
    ]
    add_table(slide, 7, 3, Inches(0.75), Inches(1.45), Inches(7.2), Inches(3.75), data)
    add_card(slide, "讲法建议", "把意图树类比为 RAG 的 Router：\n\n先决定“去哪里找”，再决定“怎么把材料交给模型”。\n\n这也是 RA 与简单向量检索 Demo 的区别。", Inches(8.25), Inches(1.65), Inches(4.3), Inches(2.15), BLUE, 15, 13)
    add_card(slide, "项目策略", "子问题并行分类，低分过滤，总意图数上限控制，避免多问句导致检索爆炸。", Inches(8.25), Inches(4.25), Inches(4.3), Inches(1.0), TEAL, 15, 13)

    # 12
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "IntentResolver 细节：阈值、并发、总量控制", "意图结果要既有召回覆盖，又不能把上下文带偏", 12)
    add_card(slide, "并发分类", "每个 subQuestion 通过 intentClassifyExecutor 并行调用分类器，降低多问句延迟。", Inches(0.78), Inches(1.45), Inches(3.8), Inches(1.35), BLUE)
    add_card(slide, "低分过滤", "INTENT_MIN_SCORE = 0.35\n\n低于阈值的意图不会进入后续路由。", Inches(4.78), Inches(1.45), Inches(3.8), Inches(1.35), TEAL)
    add_card(slide, "数量上限", "MAX_INTENT_COUNT = 3\n\n每个子问题至少保留一个，再按分数分配剩余额度。", Inches(8.78), Inches(1.45), Inches(3.8), Inches(1.35), AMBER)
    add_code_box(slide, "resolve(rewriteResult)\n  -> subQuestions\n  -> classifyIntents(q)\n  -> filter(score >= 0.35)\n  -> limit(3)\n  -> capTotalIntents(...)\n\nmergeIntentGroup(...)\n  -> kbIntents\n  -> mcpIntents", Inches(1.35), Inches(3.65), Inches(10.65), Inches(2.25), "IntentResolver")

    # 13
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "歧义引导：比误答更好的选择是先澄清", "当多个知识方向都像正确答案时，系统先问清楚，避免错误检索", 13)
    add_card(slide, "触发条件", "取 top-2 KB 意图分数：\nratio = secondScore / firstScore", Inches(0.78), Inches(1.5), Inches(3.8), Inches(1.3), BLUE)
    add_card(slide, "明确歧义", "ratio >= 0.8\n\n直接生成引导问题。", Inches(4.78), Inches(1.5), Inches(3.8), Inches(1.3), AMBER)
    add_card(slide, "灰色地带", "0.65 <= ratio < 0.8\n\n调用 LLM 二次确认。", Inches(8.78), Inches(1.5), Inches(3.8), Inches(1.3), TEAL)
    add_card(slide, "例子", "用户问：数据安全规范是什么？\n\n可能命中：OA 系统数据安全 / 互联网保险系统数据安全。\n\n直接回答可能答错业务域；先澄清能显著降低错误。", Inches(1.25), Inches(3.65), Inches(10.8), Inches(1.75), GREEN, 16, 14)

    # 14
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "RetrievalEngine：按子问题并行构建上下文", "它不是只调用向量库，而是把 KB 检索与工具结果合并成 RetrievalContext", 14)
    add_flow(slide, ["SubQuestionIntent", "KB / MCP 分组", "并行上下文构建", "格式化 evidence", "合并 RetrievalContext"], Inches(0.7), Inches(1.55), Inches(2.15), Inches(0.25), color=BLUE)
    add_card(slide, "KB 路径", "KB intents -> MultiChannelRetrievalEngine -> formatKbContext -> kbContext", Inches(0.9), Inches(3.1), Inches(3.75), Inches(1.25), BLUE, 15, 12.8)
    add_card(slide, "MCP 路径", "MCP intents -> 参数抽取 -> 工具执行 -> formatMcpContext -> mcpContext", Inches(4.85), Inches(3.1), Inches(3.75), Inches(1.25), AMBER, 15, 12.8)
    add_card(slide, "多子问题合并", "多个子问题分别包装为 sub-question-kb-wrapper / sub-question-mcp-wrapper，避免上下文混在一起。", Inches(8.8), Inches(3.1), Inches(3.75), Inches(1.25), TEAL, 15, 12.8)
    add_code_box(slide, "retrieve(subIntents, topK)\n  -> CompletableFuture per subQuestion\n  -> buildSubQuestionContext(...)\n  -> retrieveAndRerank(...)\n  -> executeMcpAndMerge(...)\n  -> RetrievalContext(kbContext, mcpContext, intentChunks)", Inches(1.4), Inches(5.1), Inches(10.55), Inches(1.05), "RetrievalEngine")

    # 15
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "多通道召回：什么时候定向，什么时候全局兜底", "这一页是本次分享重点：精准召回和兜底召回如何共存", 15)
    data = [
        ["通道", "启用条件", "作用"],
        ["IntentDirectedSearch", "配置开启；存在 KB 意图；意图分数达到 minIntentScore", "按意图绑定 collection 精准检索"],
        ["VectorGlobalSearch", "定向关闭；无意图；maxScore < 0.6；单一中等置信度", "跨全部 KB collection 兜底/补充"],
        ["PostProcessors", "检索结束后按 order 顺序执行", "去重、重排、截断"],
    ]
    add_table(slide, 4, 3, Inches(0.75), Inches(1.45), Inches(12.0), Inches(2.25), data)
    add_image_fit(slide, ASSETS / "multi-channel-retrieval.png", Inches(0.95), Inches(4.05), Inches(5.4), Inches(2.45))
    add_card(slide, "关键配置", "vector-global.confidence-threshold = 0.6\nvector-global.top-k-multiplier = 3\nintent-directed.min-intent-score = 0.4\nintent-directed.top-k-multiplier = 2", Inches(6.75), Inches(4.1), Inches(5.55), Inches(1.55), AMBER, 15, 12.5)

    # 16
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "两个检索通道内部怎么工作", "IntentDirected 负责准，VectorGlobal 负责兜底；二者并行执行，结果进入统一后处理", 16)
    add_card(slide, "IntentDirectedSearchChannel", "priority = 1\n\n1. 提取 KB 意图\n2. 按意图对应 collection 并行检索\n3. 每个意图取 topK * 2\n4. 返回 channel result", Inches(0.75), Inches(1.5), Inches(5.75), Inches(2.4), BLUE, 17, 13)
    add_card(slide, "VectorGlobalSearchChannel", "priority = 10\n\n1. 获取所有 KB collection\n2. 对全部 collection 并行检索\n3. 取 topK * 3 增大召回池\n4. 作为低置信度兜底", Inches(6.85), Inches(1.5), Inches(5.75), Inches(2.4), TEAL, 17, 13)
    add_card(slide, "异常策略", "每个通道异常时返回 emptyResult，不中断整个检索链路。\n\n这点对线上稳定性很重要：一个通道挂了，另一个通道仍可提供结果。", Inches(1.35), Inches(4.75), Inches(10.75), Inches(1.15), GREEN, 16, 13.2)

    # 17
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "后处理链：去重、重排、TopK 截断", "召回阶段追求覆盖，后处理阶段负责把上下文变干净、变有序", 17)
    add_flow(slide, ["Channel Results", "Merge Chunks", "Dedup", "Rerank", "Final TopK"], Inches(0.92), Inches(1.55), Inches(2.1), Inches(0.35), color=BLUE)
    data = [
        ["处理器", "顺序", "细节"],
        ["Deduplication", "order = 1", "按 chunk id 或 content hash 去重；按通道优先级保留，再取最高分"],
        ["Rerank", "order = 10", "调用 RerankService.rerank(query, chunks, topK)；TopK 截断发生在这里"],
        ["无 Rerank", "-", "merged + dedup 后全部流入 Prompt，可能导致上下文噪声和长度膨胀"],
    ]
    add_table(slide, 4, 3, Inches(0.85), Inches(2.75), Inches(11.8), Inches(2.2), data)
    add_card(slide, "讲述重点", "Rerank 不是锦上添花，它是“召回池”和“最终证据”之间的最后一道质量门。", Inches(1.6), Inches(5.55), Inches(10.2), Inches(0.75), AMBER, 15, 13.2)

    # 18
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Prompt 证据组装：把检索结果变成模型可用上下文", "模型能否基于证据回答，取决于证据质量，也取决于消息结构", 18)
    add_card(slide, "场景规划", "RAGPromptService.plan(context)\n\nKB_ONLY：只用知识库证据\nMCP_ONLY：只用工具结果\nMIXED：知识库 + 工具结果混合", Inches(0.78), Inches(1.45), Inches(4.0), Inches(2.0), BLUE, 15, 12.5)
    add_card(slide, "模板选择", "单意图 + 节点有 promptTemplate：使用节点模板。\n\n否则使用 answer-chat-kb / answer-chat-mcp / mixed 默认模板。", Inches(4.98), Inches(1.45), Inches(4.0), Inches(2.0), TEAL, 15, 12.5)
    add_card(slide, "消息顺序", "1. System Prompt\n2. History Messages\n3. Evidence + User Question\n\n摘要作为 history 的 system message 紧跟系统提示。", Inches(9.18), Inches(1.45), Inches(3.55), Inches(2.0), AMBER, 15, 12.5)
    add_code_box(slide, "buildStructuredMessages(context, history, question, subQuestions)\n  messages.add(systemPrompt)\n  messages.addAll(history)\n  userContent = mergeEvidenceAndQuestion(evidenceBody, userQuestion)\n  messages.add(ChatMessage.user(userContent))", Inches(1.2), Inches(4.35), Inches(10.95), Inches(1.45), "RAGPromptService")

    # 19
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Trace / Debug：RAG 链路出错时怎么查", "生产化 RAG 的关键不是永远不出错，而是出错后能定位到阶段", 19)
    add_image_fit(slide, ASSETS / "admin-trace.png", Inches(0.6), Inches(1.35), Inches(6.05), Inches(5.45))
    data = [
        ["Trace 节点", "观察什么"],
        ["query-rewrite", "改写耗时、是否失败"],
        ["intent-resolve", "意图分类耗时和状态"],
        ["guidance-detect", "是否触发歧义引导"],
        ["retrieval-engine", "检索总耗时"],
        ["multi-channel-retrieval", "通道级召回情况"],
        ["llm-stream-routing", "模型路由和流式生成"],
    ]
    add_table(slide, 7, 2, Inches(6.95), Inches(1.45), Inches(5.65), Inches(3.35), data)
    add_card(slide, "排查路径", "答案差：先看是否召回到正确文档；召回正确再看 Rerank / Prompt；召回错误再看 Rewrite / Intent / Chunking。", Inches(6.95), Inches(5.15), Inches(5.65), Inches(0.95), AMBER, 14, 12)

    # 20
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "评估闭环：用指标定位优化方向", "不要只问“回答好不好”，要拆成意图、检索、生成、性能四层看", 20)
    data = [
        ["层次", "指标", "低分时优先检查"],
        ["意图", "Top-1 Accuracy / micro-F1", "意图树、示例、分类 Prompt"],
        ["检索", "Hit@K / Recall@K / MRR", "Chunking、路由、通道阈值、向量模型"],
        ["生成", "faithfulness / answer_correctness", "Prompt 约束、证据顺序、模型参数"],
        ["工具", "Tool name + args match", "MCP 意图、参数抽取 Prompt、schema"],
        ["拒答", "误拒率 / 错答率", "空召回策略、无依据回答约束"],
        ["性能", "端到端延迟 / TTFT", "并发池、模型路由、通道数量"],
    ]
    add_table(slide, 7, 3, Inches(0.75), Inches(1.45), Inches(12.0), Inches(3.55), data)
    add_card(slide, "最后收束", "RA 项目的 RAG 工程化可以总结为三句话：\n\n可路由：先判断去哪里找。\n可扩展：检索通道、后处理器、入库节点都能插拔。\n可评估：Trace + Eval + Feedback 让优化可复盘。", Inches(1.05), Inches(5.45), Inches(11.4), Inches(1.15), GREEN, 15, 13)

    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    print(make_detailed_deck())
