from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "assets"
OUT = ROOT / "docs" / "RA-RAG-组会汇报.pptx"


WIDE_W = Inches(13.333)
WIDE_H = Inches(7.5)

BG = RGBColor(248, 250, 252)
INK = RGBColor(20, 33, 61)
MUTED = RGBColor(86, 100, 120)
BLUE = RGBColor(37, 99, 235)
TEAL = RGBColor(13, 148, 136)
AMBER = RGBColor(245, 158, 11)
RED = RGBColor(220, 38, 38)
GREEN = RGBColor(22, 163, 74)
LINE = RGBColor(218, 226, 238)
WHITE = RGBColor(255, 255, 255)
DARK = RGBColor(15, 23, 42)


def set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = LINE


def set_text_frame(tf, font_size=18, color=INK, bold=False):
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    for p in tf.paragraphs:
        for run in p.runs:
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = color


def add_bg(slide, color=BG):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, WIDE_W, WIDE_H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_footer(slide, idx):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(7.08), Inches(12.25), Inches(0.22))
    tf = box.text_frame
    tf.text = f"RA 项目 RAG 工程化实践  |  {idx:02d}"
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p.runs[0].font.name = "Microsoft YaHei"
    p.runs[0].font.size = Pt(8.5)
    p.runs[0].font.color.rgb = RGBColor(125, 139, 160)


def add_title(slide, title, subtitle=None, idx=None):
    box = slide.shapes.add_textbox(Inches(0.65), Inches(0.38), Inches(8.8), Inches(0.62))
    tf = box.text_frame
    tf.text = title
    p = tf.paragraphs[0]
    p.runs[0].font.name = "Microsoft YaHei"
    p.runs[0].font.size = Pt(27)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = INK
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.68), Inches(1.02), Inches(9.6), Inches(0.34))
        stf = sub.text_frame
        stf.text = subtitle
        stf.paragraphs[0].runs[0].font.name = "Microsoft YaHei"
        stf.paragraphs[0].runs[0].font.size = Pt(11.5)
        stf.paragraphs[0].runs[0].font.color.rgb = MUTED
    if idx is not None:
        add_footer(slide, idx)


def add_badge(slide, text, x, y, color=BLUE):
    badge = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, Inches(1.3), Inches(0.34))
    badge.fill.solid()
    badge.fill.fore_color.rgb = color
    badge.line.fill.background()
    tf = badge.text_frame
    tf.text = text
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.runs[0].font.name = "Microsoft YaHei"
    p.runs[0].font.size = Pt(10)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = WHITE


def add_bullets(slide, items, x, y, w, h, font_size=17, color=INK, bullet_color=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(7)
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
    return box


def add_card(slide, title, body, x, y, w, h, accent=BLUE, title_size=15, body_size=12.5):
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    set_fill(card, WHITE)
    card.line.color.rgb = LINE
    card.shadow.inherit = False
    stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, Inches(0.08), h)
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()
    title_box = slide.shapes.add_textbox(x + Inches(0.22), y + Inches(0.16), w - Inches(0.38), Inches(0.32))
    title_box.text_frame.text = title
    p = title_box.text_frame.paragraphs[0]
    p.runs[0].font.name = "Microsoft YaHei"
    p.runs[0].font.size = Pt(title_size)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = INK
    body_box = slide.shapes.add_textbox(x + Inches(0.22), y + Inches(0.57), w - Inches(0.38), h - Inches(0.66))
    tf = body_box.text_frame
    tf.text = body
    tf.word_wrap = True
    for p in tf.paragraphs:
        for run in p.runs:
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(body_size)
            run.font.color.rgb = MUTED
    return card


def add_image_fit(slide, image_path, x, y, w, h, border=True):
    image_path = Path(image_path)
    with Image.open(image_path) as img:
        iw, ih = img.size
    scale = min(w / iw, h / ih)
    pic_w = int(iw * scale)
    pic_h = int(ih * scale)
    px = x + int((w - pic_w) / 2)
    py = y + int((h - pic_h) / 2)
    if border:
        frame = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
        set_fill(frame, WHITE)
        frame.line.color.rgb = LINE
    slide.shapes.add_picture(str(image_path), px, py, width=pic_w, height=pic_h)


def add_section_label(slide, text, x, y, color=TEAL):
    box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, Inches(2.1), Inches(0.38))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(231, 248, 246)
    box.line.color.rgb = RGBColor(167, 224, 218)
    tf = box.text_frame
    tf.text = text
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.runs[0].font.name = "Microsoft YaHei"
    p.runs[0].font.size = Pt(10)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = color


def connect(slide, x1, y1, x2, y2, color=LINE):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    line.line.color.rgb = color
    line.line.width = Pt(1.8)
    return line


def add_flow(slide, labels, x, y, box_w, gap, color=BLUE):
    for i, label in enumerate(labels):
        bx = x + i * (box_w + gap)
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, bx, y, box_w, Inches(0.62))
        set_fill(shape, WHITE)
        shape.line.color.rgb = color if i == 0 else LINE
        tf = shape.text_frame
        tf.text = label
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.runs[0].font.name = "Microsoft YaHei"
        p.runs[0].font.size = Pt(11.2)
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = INK
        if i < len(labels) - 1:
            connect(slide, bx + box_w, y + Inches(0.31), bx + box_w + gap, y + Inches(0.31), color=RGBColor(160, 174, 194))


def make_deck():
    prs = Presentation()
    prs.slide_width = WIDE_W
    prs.slide_height = WIDE_H
    blank = prs.slide_layouts[6]

    # 1
    slide = prs.slides.add_slide(blank)
    add_bg(slide, DARK)
    slide.shapes.add_picture(str(ASSETS / "ragent-ai-banner.png"), Inches(8.25), Inches(0.35), width=Inches(4.3))
    add_badge(slide, "组会分享", Inches(0.75), Inches(0.82), TEAL)
    title = slide.shapes.add_textbox(Inches(0.75), Inches(1.52), Inches(8.5), Inches(1.6))
    title.text_frame.text = "从 RA 项目看 RAG 系统的工程化实践"
    p = title.text_frame.paragraphs[0]
    p.runs[0].font.name = "Microsoft YaHei"
    p.runs[0].font.size = Pt(36)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = WHITE
    sub = slide.shapes.add_textbox(Inches(0.78), Inches(3.15), Inches(8.3), Inches(0.68))
    sub.text_frame.text = "从文档入库、多路召回到可观测评估的完整链路"
    p = sub.text_frame.paragraphs[0]
    p.runs[0].font.name = "Microsoft YaHei"
    p.runs[0].font.size = Pt(18)
    p.runs[0].font.color.rgb = RGBColor(203, 213, 225)
    add_flow(slide, ["离线入库", "在线链路", "多路召回", "可观测评估"], Inches(0.82), Inches(5.1), Inches(2.15), Inches(0.35), color=TEAL)
    add_footer(slide, 1)

    # 2
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "15-20 分钟汇报路线", "不逐题讲概念，而是把 RAG 知识点落到 RA 项目链路上", 2)
    cols = [
        ("1. 先看地图", "系统全景\nRAG 解决什么问题", BLUE),
        ("2. 再看入库", "文档解析\n结构化 Chunking\n向量化与索引", TEAL),
        ("3. 重点讲在线链路", "Query Rewrite\n意图路由\n歧义引导\n多路召回", AMBER),
        ("4. 最后讲落地", "Rerank 与 Prompt\nTrace 定位\n评估闭环", GREEN),
    ]
    for i, (t, b, c) in enumerate(cols):
        add_card(slide, t, b, Inches(0.75 + i * 3.1), Inches(2.0), Inches(2.72), Inches(3.35), c, 15, 15)
    add_section_label(slide, "砍掉：模块分层 / Embedding 细节 / MCP 专题", Inches(0.75), Inches(6.05), RED)

    # 3
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "RA 的 RAG 系统全景", "一次回答不是“向量检索 + LLM”，而是一条可编排、可短路、可观测的工程链路", 3)
    add_image_fit(slide, ASSETS / "ragent-chain-v3.png", Inches(0.68), Inches(1.52), Inches(7.45), Inches(5.1))
    add_card(slide, "讲述重点", "把听众先带到同一张地图上：\n\n用户问题进入系统后，会依次经历记忆加载、问题改写、意图解析、歧义处理、检索、Prompt 组装和流式生成。", Inches(8.35), Inches(1.65), Inches(4.25), Inches(2.15), BLUE, 15, 13.2)
    add_card(slide, "核心判断", "RAG 工程化的难点不在“能不能召回”，而在：\n\n召回哪里、召回多少、如何筛选、如何证明回答可信。", Inches(8.35), Inches(4.05), Inches(4.25), Inches(1.85), TEAL, 15, 13.2)

    # 4
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "为什么 RA 项目需要 RAG", "RAG 的目标不是替代模型能力，而是把私域知识变成可控、可追溯的回答依据", 4)
    add_card(slide, "知识时效性", "业务文档、制度、产品说明持续变化，RAG 可以通过知识库更新吸收新内容。", Inches(0.75), Inches(1.72), Inches(3.7), Inches(1.35), BLUE)
    add_card(slide, "私域知识", "模型预训练里没有企业内部流程、业务规则、系统文档，需要从项目知识库检索。", Inches(4.82), Inches(1.72), Inches(3.7), Inches(1.35), TEAL)
    add_card(slide, "可追溯", "回答可以绑定检索证据，后续能通过 Trace、评估和反馈定位问题。", Inches(8.9), Inches(1.72), Inches(3.7), Inches(1.35), GREEN)
    add_flow(slide, ["用户问题", "检索证据", "约束生成", "可评估答案"], Inches(1.4), Inches(4.55), Inches(2.3), Inches(0.55), color=BLUE)
    add_bullets(slide, ["这一页只讲项目动机，不展开微调对比", "RAG 更适合外部知识注入、持续更新和答案可追溯"], Inches(1.15), Inches(5.65), Inches(10.8), Inches(0.75), 15, MUTED)

    # 5
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "离线入库：文档如何变成可检索知识", "入库链路决定了召回上限；在线检索很难弥补低质量 Chunk", 5)
    add_image_fit(slide, ASSETS / "ingestion-pipeline.png", Inches(0.65), Inches(1.42), Inches(6.6), Inches(5.25))
    steps = [
        ("Fetcher", "拉取本地、远程、S3/Feishu 等来源"),
        ("Parser", "解析 Markdown、Excel、PDF、图片等格式"),
        ("Chunker", "按结构与语义切分，生成 VectorChunk"),
        ("Indexer", "Embedding 后写入 PgVector / Milvus"),
    ]
    for i, (t, b) in enumerate(steps):
        add_card(slide, t, b, Inches(7.55), Inches(1.45 + i * 1.22), Inches(4.8), Inches(0.92), [BLUE, TEAL, AMBER, GREEN][i], 14, 11.5)

    # 6
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Chunking：切块不是按字数切完就结束", "结构化分块的目标：保留语义边界，同时控制每个 Chunk 的上下文体量", 6)
    add_card(slide, "优先走结构化 Block", "Parser 如果产出 heading / paragraph / table / list / image / code 等结构，Chunker 会按 block-aware 策略处理。", Inches(0.75), Inches(1.58), Inches(5.75), Inches(1.35), BLUE)
    add_card(slide, "没有结构时走文本策略", "blocks 为空时，系统回退到固定大小或语义感知文本切分，并保留 overlap。", Inches(0.75), Inches(3.15), Inches(5.75), Inches(1.35), TEAL)
    add_card(slide, "特殊内容单独处理", "表格控制行数，列表控制 item 数，图片先转知识文本，避免重要信息被切散。", Inches(0.75), Inches(4.72), Inches(5.75), Inches(1.35), AMBER)
    add_card(slide, "讲清一个取舍", "Chunk 太大：召回噪声高、上下文浪费。\nChunk 太小：语义被切断、context recall 下降。", Inches(7.05), Inches(1.72), Inches(5.35), Inches(3.05), RED, 16, 15)
    add_section_label(slide, "建议讲 1.5 分钟即可", Inches(8.48), Inches(5.35), TEAL)

    # 7
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "在线链路：一次问题如何走完整个系统", "这页是汇报主轴，后面几页都围绕这条 Pipeline 展开", 7)
    labels = ["loadMemory", "rewriteQuery", "resolveIntents", "guidance", "retrieve", "streamRagResponse"]
    add_flow(slide, labels, Inches(0.72), Inches(1.65), Inches(1.65), Inches(0.25), color=BLUE)
    add_card(slide, "短路分支", "歧义问题：先返回引导问题\n系统闲聊：不做检索直接回复\n无检索结果：明确提示没有相关文档", Inches(0.85), Inches(3.05), Inches(3.75), Inches(2.18), AMBER, 15, 13.2)
    add_card(slide, "工程价值", "把 RAG 拆成可观测阶段，而不是一个黑盒函数。\n\n每个阶段都能单独调参、替换、打 Trace。", Inches(4.95), Inches(3.05), Inches(3.75), Inches(2.18), TEAL, 15, 13.2)
    add_card(slide, "代码锚点", "StreamChatPipeline.execute()\n\n这条链路定义了 RA 项目的在线问答主流程。", Inches(9.05), Inches(3.05), Inches(3.5), Inches(2.18), BLUE, 15, 13.2)

    # 8
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Query Rewrite：把用户问法转换成检索友好的问法", "它解决的是“用户怎么问”和“知识怎么存”之间的不一致", 8)
    add_card(slide, "做什么", "结合会话历史，把口语化、指代、省略、多问句整理成更适合检索的查询。", Inches(0.75), Inches(1.55), Inches(3.8), Inches(1.45), BLUE)
    add_card(slide, "为什么", "向量检索对 query 表达很敏感；改写可以提高召回覆盖，拆分可以减少多意图互相干扰。", Inches(4.82), Inches(1.55), Inches(3.8), Inches(1.45), TEAL)
    add_card(slide, "失败怎么办", "改写失败时回退原问题；多问题拆分失败时可以规则兜底，避免链路中断。", Inches(8.9), Inches(1.55), Inches(3.8), Inches(1.45), AMBER)
    add_card(slide, "例子", "原问题：报销怎么弄，发票要啥？\n\n拆分后：\n1. 报销流程是什么？\n2. 开票/发票信息需要哪些材料？", Inches(1.45), Inches(4.0), Inches(10.45), Inches(1.55), GREEN, 16, 15)

    # 9
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "意图路由：先判断该去哪里找", "意图识别决定检索目标，也决定后续 Prompt 场景", 9)
    add_card(slide, "KB", "进入知识库检索：制度、文档、FAQ、产品说明。", Inches(0.85), Inches(1.7), Inches(3.5), Inches(1.45), BLUE, 17, 14)
    add_card(slide, "SYSTEM", "纯系统交互：打招呼、自我介绍、无需检索的直接回复。", Inches(4.95), Inches(1.7), Inches(3.5), Inches(1.45), TEAL, 17, 14)
    add_card(slide, "MCP", "工具调用场景：实时数据、外部动作、结构化查询。这里只点到，不做专题展开。", Inches(9.05), Inches(1.7), Inches(3.5), Inches(1.45), AMBER, 17, 14)
    add_card(slide, "项目策略", "对每个子问题并行分类，过滤低分意图，并限制总意图数量。\n\n这个策略让检索既能聚焦，又不会因为多问句失控。", Inches(1.15), Inches(4.05), Inches(11.0), Inches(1.55), GREEN, 16, 14)

    # 10
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "歧义引导：不确定时先问清楚", "比起盲目召回，先澄清问题能显著降低错误答案风险", 10)
    add_card(slide, "触发场景", "同一个问题命中多个相近 KB 意图，且 top-2 分数接近。", Inches(0.85), Inches(1.6), Inches(3.55), Inches(1.45), AMBER)
    add_card(slide, "系统动作", "短路正常检索流程，直接返回一个引导用户选择的澄清问题。", Inches(4.95), Inches(1.6), Inches(3.55), Inches(1.45), BLUE)
    add_card(slide, "价值", "减少误路由、误召回和幻觉，让系统在不确定时表现得更稳。", Inches(9.05), Inches(1.6), Inches(3.55), Inches(1.45), GREEN)
    add_card(slide, "可以这样举例", "用户问：数据安全规范是什么？\n\n系统可引导：你想了解 OA 系统的数据安全，还是互联网保险系统的数据安全？", Inches(1.2), Inches(4.05), Inches(10.9), Inches(1.6), TEAL, 16, 15)

    # 11
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "多路召回：RA 项目的检索核心", "高置信度走定向检索，低置信度用全局检索兜底或补充", 11)
    add_image_fit(slide, ASSETS / "multi-channel-retrieval.png", Inches(0.72), Inches(1.35), Inches(6.45), Inches(5.45))
    add_card(slide, "IntentDirectedSearch", "优先级 1。\n\n识别到 KB 意图后，去意图绑定的 collection 检索。适合明确问题，精准度高。", Inches(7.45), Inches(1.55), Inches(4.85), Inches(1.55), BLUE, 15, 13)
    add_card(slide, "VectorGlobalSearch", "优先级 10。\n\n当意图低置信度、定向通道关闭或单一中等置信度时启用，负责兜底和补充。", Inches(7.45), Inches(3.35), Inches(4.85), Inches(1.55), TEAL, 15, 13)
    add_card(slide, "讲述关键词", "并行、互补、阈值控制、TopK multiplier、失败不拖垮整体链路。", Inches(7.45), Inches(5.15), Inches(4.85), Inches(0.92), AMBER, 15, 13)

    # 12
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Rerank 后处理：召回多不等于上下文好", "后处理链把多通道结果变成最终喂给 LLM 的高质量证据", 12)
    add_flow(slide, ["多通道原始结果", "Dedup 去重", "Rerank 重排", "TopK 截断", "最终上下文"], Inches(0.85), Inches(1.7), Inches(2.05), Inches(0.3), color=BLUE)
    add_card(slide, "Deduplication", "始终启用，按 Chunk ID 或内容 hash 去重；同一个 Chunk 多通道命中时保留更高分结果。", Inches(0.95), Inches(3.0), Inches(5.35), Inches(1.45), TEAL, 16, 13.5)
    add_card(slide, "Rerank", "最后执行，调用 Rerank 模型重新排序；项目里 TopK 截断也发生在这里。", Inches(7.0), Inches(3.0), Inches(5.35), Inches(1.45), AMBER, 16, 13.5)
    add_card(slide, "要强调的坑", "如果关闭 Rerank，只剩 merged + dedup 的结果进入生成，可能出现上下文过多或噪声变大。", Inches(1.85), Inches(5.25), Inches(9.8), Inches(0.85), RED, 15, 13.2)

    # 13
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Prompt 证据组装：让模型基于材料回答", "Prompt 的重点不是写得花，而是把系统指令、历史、证据和问题放在正确位置", 13)
    add_card(slide, "场景选择", "KB_ONLY / MCP_ONLY / MIXED 三类场景选择不同模板。", Inches(0.75), Inches(1.55), Inches(3.85), Inches(1.35), BLUE)
    add_card(slide, "消息顺序", "system prompt → conversation history → evidence + question。", Inches(4.75), Inches(1.55), Inches(3.85), Inches(1.35), TEAL)
    add_card(slide, "证据格式", "KB context 和工具数据会被格式化成结构化上下文，再与用户问题合并。", Inches(8.75), Inches(1.55), Inches(3.85), Inches(1.35), GREEN)
    add_card(slide, "讲述重点", "减少幻觉不是只靠一句“不要编造”。\n\n更关键的是：检索证据质量、上下文顺序、Prompt 场景、温度参数共同约束模型输出。", Inches(1.35), Inches(4.0), Inches(10.6), Inches(1.55), AMBER, 16, 15)

    # 14
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Trace 与评估：RAG 出错时怎么定位", "组会里这页很重要：证明系统不是黑盒，而是能诊断、能复盘、能迭代", 14)
    add_image_fit(slide, ASSETS / "admin-trace.png", Inches(0.68), Inches(1.42), Inches(6.25), Inches(5.25))
    add_card(slide, "Trace 定位", "关键节点打 @RagTraceNode：rewrite、intent、retrieval、multi-channel、LLM 等阶段都能看耗时和状态。", Inches(7.25), Inches(1.45), Inches(5.05), Inches(1.3), BLUE, 15, 12.5)
    add_card(slide, "评估指标", "意图准确率、Hit@K / Recall@K / MRR、Tool 调用准确率、拒答正确性、RAGAS 五指标。", Inches(7.25), Inches(3.05), Inches(5.05), Inches(1.3), TEAL, 15, 12.5)
    add_card(slide, "排障读法", "召回漏了看 chunk / 路由；召回脏了看 rerank / topK；答案编了看 Prompt / 模型约束。", Inches(7.25), Inches(4.65), Inches(5.05), Inches(1.3), AMBER, 15, 12.5)

    # 15
    slide = prs.slides.add_slide(blank)
    add_bg(slide, DARK)
    title = slide.shapes.add_textbox(Inches(0.75), Inches(0.65), Inches(7.6), Inches(0.72))
    title.text_frame.text = "总结：RAG 工程化的三个关键词"
    p = title.text_frame.paragraphs[0]
    p.runs[0].font.name = "Microsoft YaHei"
    p.runs[0].font.size = Pt(31)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = WHITE
    cards = [
        ("可路由", "Query Rewrite + 意图树，让问题先找到正确方向。", BLUE),
        ("可扩展", "SearchChannel / PostProcessor / IngestionNode 都是扩展点。", TEAL),
        ("可评估", "Trace + Feedback + Eval，让优化有依据。", AMBER),
    ]
    for i, (t, b, c) in enumerate(cards):
        add_card(slide, t, b, Inches(0.95 + i * 4.05), Inches(2.1), Inches(3.45), Inches(2.35), c, 22, 16)
    closing = slide.shapes.add_textbox(Inches(1.25), Inches(5.35), Inches(10.8), Inches(0.6))
    closing.text_frame.text = "这次分享的核心：不是介绍一个 RAG Demo，而是拆解一套能上线、能调优、能复盘的 RAG 链路。"
    p = closing.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.runs[0].font.name = "Microsoft YaHei"
    p.runs[0].font.size = Pt(18)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor(226, 232, 240)
    add_footer(slide, 15)

    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = make_deck()
    print(path)
